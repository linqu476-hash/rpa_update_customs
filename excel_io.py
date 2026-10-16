"""
Excel template + reader for the Customs No / Invoice No upload.

    build_template()  -> bytes of the .xlsx template
    read_invoice_file(data: bytes, filename) -> dict(groups, rows, errors, warnings)
"""
import io
import re

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FONT = "Arial"
HEADERS = ["Customs No", "Invoice No"]

# header names accepted when reading (compared after removing spaces/punctuation, lower case)
CUSTOMS_NAMES = {"customsno", "customsnumber", "customs", "custom", "customno"}
INVOICE_NAMES = {"invoiceno", "invoicenumber", "invoice", "invno", "inv", "exportinvoiceid", "exportinvoiceno"}


def build_template(rows=None) -> bytes:
    """Empty template, or one prefilled with rows [(customs_no, invoice_no), ...]."""
    wb = Workbook()

    # ---- Sheet 1: Invoices (the sheet that gets uploaded) ----
    ws = wb.active
    ws.title = "Invoices"
    head_fill = PatternFill("solid", fgColor="17415F")
    input_fill = PatternFill("solid", fgColor="FFF9DB")
    thin = Side(style="thin", color="C9D2DA")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, name in enumerate(HEADERS, start=1):
        c = ws.cell(row=1, column=col, value=name)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=11)
        c.fill = head_fill
        c.alignment = Alignment(vertical="center")
        c.border = border
    ws["A1"].comment = Comment("Customs number, for example E274120. Leave blank to repeat the number above.", "Template")
    ws["B1"].comment = Comment("Invoice number exactly as in D365, for example C265825-1.", "Template")

    for r in range(2, 202):                     # 200 input rows, formatted as text so '00123' stays '00123'
        for col in range(1, 3):
            c = ws.cell(row=r, column=col)
            c.number_format = "@"
            c.font = Font(name=FONT, size=11)
            c.border = border
            if r <= 51:
                c.fill = input_fill
    for i, row in enumerate(rows or [], start=2):
        for col, v in enumerate(row, start=1):
            c = ws.cell(row=i, column=col, value=v)
            c.number_format = "@"
            c.font = Font(name=FONT, size=11)
            c.border = border
            if i <= 51:
                c.fill = input_fill
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 24
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"

    # ---- Sheet 2: Instructions ----
    ins = wb.create_sheet("Instructions")
    ins.sheet_view.showGridLines = False
    ins.column_dimensions["A"].width = 18
    ins.column_dimensions["B"].width = 24
    ins.column_dimensions["C"].width = 50

    ins["A1"] = "How to fill this template"
    ins["A1"].font = Font(name=FONT, bold=True, size=14)
    lines = [
        "1. Fill the sheet 'Invoices' only. The shaded cells (rows 2 to 51) are for your data; you can use more rows below.",
        "2. One row per invoice. Customs No can be left blank to repeat the number from the row above.",
        "3. Do not rename the headers. Do not leave the example rows below in the 'Invoices' sheet.",
        "4. Save the file, then upload it in the dashboard under 'Update Customs No'.",
    ]
    for i, text in enumerate(lines, start=3):
        ins.cell(row=i, column=1, value=text).font = Font(name=FONT, size=11)

    ins["A9"] = "Example of filled rows (illustration only)"
    ins["A9"].font = Font(name=FONT, bold=True, size=11)
    for col, name in enumerate(HEADERS, start=1):
        c = ins.cell(row=10, column=col, value=name)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = head_fill
        c.border = border
    examples = [("E274120", "C265825"), ("", "C265825-1"), ("E274121", "C265240")]
    for r, row in enumerate(examples, start=11):
        for col, v in enumerate(row, start=1):
            c = ins.cell(row=r, column=col, value=v)
            c.font = Font(name=FONT, size=11, color="0000FF")   # blue = typed input
            c.border = border
    ins["C11"] = "Customs No E274120 with its first invoice"
    ins["C12"] = "Blank Customs No: same as the row above (E274120)"
    ins["C13"] = "A new Customs No starts a new group"
    for r in (11, 12, 13):
        ins.cell(row=r, column=3).font = Font(name=FONT, size=10, color="5B6B7A")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _norm(v) -> str:
    return re.sub(r"[^a-z0-9]", "", str(v).lower())


def _clean(v) -> str:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    s = str(v).strip()
    if s.lower() in ("nan", "none", "nat"):
        return ""
    if re.fullmatch(r"\d+\.0", s):          # Excel stored a number: 123 -> '123.0'
        s = s[:-2]
    return s


def read_invoice_file(data: bytes, filename: str) -> dict:
    """Read the uploaded template. Returns {groups, rows, errors, warnings}."""
    out = {"groups": {}, "rows": pd.DataFrame(columns=["Excel row", "Customs No", "Invoice No"]),
           "errors": [], "warnings": []}
    try:
        if filename.lower().endswith(".csv"):
            raw = pd.read_csv(io.BytesIO(data), header=None, dtype=str, keep_default_na=False,
                              sep=None, engine="python")
        else:
            xl = pd.ExcelFile(io.BytesIO(data))
            sheet = "Invoices" if "Invoices" in xl.sheet_names else xl.sheet_names[0]
            raw = xl.parse(sheet, header=None, dtype=str)
            if sheet != "Invoices":
                out["warnings"].append(f"No sheet named 'Invoices' - read the first sheet ('{sheet}').")
    except Exception as e:
        out["errors"].append(f"The file could not be read: {e}")
        return out

    # find the header row (first 15 rows)
    header_row = cust_col = inv_col = None
    for i in range(min(15, len(raw))):
        names = [_norm(v) for v in raw.iloc[i].tolist()]
        c = next((j for j, n in enumerate(names) if n in CUSTOMS_NAMES), None)
        v = next((j for j, n in enumerate(names) if n in INVOICE_NAMES), None)
        if c is not None and v is not None:
            header_row, cust_col, inv_col = i, c, v
            break
    if header_row is None:
        out["errors"].append("Could not find the columns 'Customs No' and 'Invoice No' in the first rows. "
                             "Use the template and keep its header names.")
        return out

    rows, last_customs = [], ""
    seen = set()
    inv_owner = {}
    for i in range(header_row + 1, len(raw)):
        cust = _clean(raw.iat[i, cust_col])
        inv = _clean(raw.iat[i, inv_col])
        excel_row = i + 1
        if not cust and not inv:
            continue
        if cust:
            last_customs = cust
        elif inv:
            cust = last_customs                   # blank Customs No = same as the row above
        if inv and not cust:
            out["errors"].append(f"Row {excel_row}: invoice {inv} has no Customs No.")
            continue
        if cust and not inv:
            out["warnings"].append(f"Row {excel_row}: Customs No {cust} has no invoice - ignored.")
            continue
        if re.search(r"\s", inv) or re.search(r"\s", cust):
            out["warnings"].append(f"Row {excel_row}: spaces inside '{cust}' / '{inv}' - check the values.")
        if (cust, inv) in seen:
            out["warnings"].append(f"Row {excel_row}: {inv} is listed twice for {cust} - kept once.")
            continue
        if inv in inv_owner and inv_owner[inv] != cust:
            out["warnings"].append(f"Row {excel_row}: {inv} appears under {inv_owner[inv]} and {cust}.")
        seen.add((cust, inv))
        inv_owner.setdefault(inv, cust)
        rows.append({"Excel row": excel_row, "Customs No": cust, "Invoice No": inv})

    if not rows and not out["errors"]:
        out["errors"].append("The file has no invoice rows. Fill the 'Invoices' sheet and upload it again.")
    out["rows"] = pd.DataFrame(rows, columns=["Excel row", "Customs No", "Invoice No"])
    groups = {}
    for r in rows:
        groups.setdefault(r["Customs No"], []).append(r["Invoice No"])
    out["groups"] = groups
    return out


def sc_base_invoice(inv: str) -> str:
    """SC260389-1, SC260389-22 ... -> SC260389. Invoices that do not start with SC are returned unchanged."""
    if re.match(r"(?i)SC\d", inv):          # SC followed by digits, e.g. SC260389-4
        return re.sub(r"(-\d+)+$", "", inv)
    return inv


def convert_raw(data: bytes, filename: str, sc_base_only: bool = True) -> dict:
    """Turn the raw customs list (columns like DATE, ETD, CUSTOM, PAGE DATE, INV#) into the template.

    CUSTOM -> Customs No, INV# -> Invoice No (other columns, such as dates, are ignored).
    Rows without an invoice (title/ETD rows, blank separators) are skipped. A blank CUSTOM cell is only
    filled from the row directly above (never across a blank separator row).
    sc_base_only: invoices starting with SC that carry a suffix (-1, -2, -3 ...) are reduced to the base
    invoice (SC260389-5 -> SC260389) and listed once. Other invoices (for example C265386-1) are kept as is.
    Returns {template (bytes|None), rows (DataFrame), errors, warnings, info}."""
    out = {"template": None, "errors": [], "warnings": [], "info": [],
           "rows": pd.DataFrame(columns=["Sheet", "Customs No", "Invoice No"])}
    try:
        if filename.lower().endswith(".csv"):
            sheets = {"CSV": pd.read_csv(io.BytesIO(data), header=None, dtype=object, keep_default_na=False,
                                         sep=None, engine="python")}
        else:
            xl = pd.ExcelFile(io.BytesIO(data))
            sheets = {n: xl.parse(n, header=None, dtype=object) for n in xl.sheet_names}
    except Exception as e:
        out["errors"].append(f"The file could not be read: {e}")
        return out

    rows, seen = [], set()
    collapsed = {}                         # base invoice -> [suffix rows that were reduced to it]
    for sheet, raw in sheets.items():
        header_row = cust_col = inv_col = None
        for i in range(min(15, len(raw))):
            names = [_norm(v) for v in raw.iloc[i].tolist()]
            c = next((j for j, n in enumerate(names) if n in CUSTOMS_NAMES), None)
            v = next((j for j, n in enumerate(names) if n in INVOICE_NAMES), None)
            if c is not None and v is not None:
                header_row, cust_col, inv_col = i, c, v
                break
        if header_row is None:
            out["warnings"].append(f"Sheet '{sheet.strip()}' skipped: no CUSTOM and INV# columns found.")
            continue

        last_cust, gap = "", True
        for i in range(header_row + 1, len(raw)):
            cust = _clean(raw.iat[i, cust_col])
            inv = _clean(raw.iat[i, inv_col])
            excel_row = i + 1
            if not cust and not inv:
                gap, last_cust = True, ""            # blank separator row: groups do not continue across it
                continue
            if not inv:
                out["warnings"].append(f"Sheet '{sheet.strip()}' row {excel_row}: Customs No {cust} has no invoice - skipped.")
                gap, last_cust = False, cust
                continue
            if not cust:
                if last_cust and not gap:
                    cust = last_cust
                else:
                    out["warnings"].append(f"Sheet '{sheet.strip()}' row {excel_row}: invoice {inv} has no Customs No - not included.")
                    continue
            last_cust, gap = cust, False
            was_collapsed = False
            if sc_base_only:
                base = sc_base_invoice(inv)
                if base != inv:
                    collapsed.setdefault((cust, base), []).append(inv)
                    inv, was_collapsed = base, True
            if (cust, inv) in seen:
                if was_collapsed:
                    continue                       # expected: the base invoice is already listed
                out["warnings"].append(f"Sheet '{sheet.strip()}' row {excel_row}: {inv} is listed twice for {cust} - kept once.")
                continue
            seen.add((cust, inv))
            rows.append({"Sheet": sheet.strip(), "Customs No": cust, "Invoice No": inv})

    if not rows:
        if not out["errors"]:
            out["errors"].append("No customs / invoice rows were found in this file.")
        return out
    for (cust, base), subs in collapsed.items():
        suffixes = [x[len(base):] for x in subs]
        shown = suffixes[0] if len(suffixes) == 1 else f"{suffixes[0]} to {suffixes[-1]}"
        out["info"].append(f"{base} ({cust}): sub-invoices {shown} ({len(subs)} rows) are not listed; only {base} is kept.")
    out["rows"] = pd.DataFrame(rows)
    out["template"] = build_template([(r["Customs No"], r["Invoice No"]) for r in rows])
    return out
